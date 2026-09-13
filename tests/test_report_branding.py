"""Phase 200 Plan 03 — cross-surface report branding tests (RPT-01).

Presence-based tests (per feedback_report_render_tests_presence_not_appearance):
assert the six report.branding fields are PRESENT in HTML and DOCX output, not
that the visual layout / cover placement matches. Visual fidelity is a human-UAT
concern (Series 200); this file only proves per-field conditionality, the
report.branding -> assessment.logo_path -> None logo precedence, and graceful
degradation on an unreadable logo. No PDF assertions — PDF is a Playwright
print of the same HTML this file already exercises.

Node IDs:
  test_html_full_branding_all_fields_present
  test_docx_full_branding_all_fields_present
  test_html_partial_branding_only_set_field_present_absent_fields_missing
  test_absent_report_section_renders_identically_to_captured_baseline
  test_logo_precedence_report_branding_wins_over_assessment
  test_logo_precedence_assessment_logo_used_alone
  test_degradation_nonexistent_logo_path_no_crash_either_surface
  test_html_escaping_of_html_special_branding_value
"""
from __future__ import annotations

import struct
import zlib
from types import SimpleNamespace


def _make_minimal_cfg(report=None):
    """Minimal cfg SimpleNamespace mirroring test_report_render_parity.py pattern.

    `report` is pinned explicitly (keyword, default None) on every fixture cfg
    built by this helper — never left unset on a spec-mocked cfg (the
    auto-vivify trap, 200-RESEARCH.md Pitfall 6).
    """
    return SimpleNamespace(
        assessment=SimpleNamespace(
            name="Test Org",
            report_owner="Test Owner",
            data_classification="CONFIDENTIAL",
            timezone="UTC",
            logo_path=None,
        ),
        output=SimpleNamespace(directory="/tmp/quirk_test_report_branding"),
        report=report,
    )


def _branding_ns(**overrides):
    fields = dict(
        logo_path=None,
        client_name=None,
        engagement_name=None,
        prepared_by=None,
        cover_date=None,
        confidentiality_line=None,
    )
    fields.update(overrides)
    return SimpleNamespace(**fields)


def _report_cfg(**branding_overrides):
    return SimpleNamespace(template_dir=None, branding=_branding_ns(**branding_overrides))


def _make_exec_content():
    from quirk.reports.content_model import ExecContent

    return ExecContent(
        narrative_lead="Test narrative lead.",
        narrative_drivers=[],
        top_risks=[],
        roadmap_items=[],
        score_total=80,
        score_band="GOOD",
        subscores={},
        raw_sum=0,
        sev_counts={},
        hardware_devices=[],
    )


def _png_bytes(r: int, g: int, b: int) -> bytes:
    """Build a minimal, valid 1x1 PNG with a distinct pixel color.

    No dependency on a real logo asset (tmp_path-generated per RPT-01 plan
    instruction) — python-docx's built-in PNG reader (no PIL needed) accepts
    this fine, verified interactively against docx.shared.Inches.add_picture.
    """
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
    idat = chunk(b"IDAT", zlib.compress(bytes([0, r, g, b])))
    iend = chunk(b"IEND", b"")
    return sig + ihdr + idat + iend


# ---------------------------------------------------------------------------
# Full branding — all six fields present on both surfaces
# ---------------------------------------------------------------------------


def test_html_full_branding_all_fields_present(tmp_path):
    from quirk.reports.html_renderer import render_html_report

    logo_path = tmp_path / "logo.png"
    logo_path.write_bytes(_png_bytes(255, 0, 0))

    cfg = _make_minimal_cfg(report=_report_cfg(
        logo_path=str(logo_path),
        client_name="Acme Client",
        engagement_name="Q3 2026 Readiness Assessment",
        prepared_by="Jane Consultant",
        cover_date="2026-09-11",
        confidentiality_line="STRICTLY CONFIDENTIAL — DO NOT DISTRIBUTE",
    ))

    path = str(tmp_path / "report.html")
    render_html_report(
        path=path, cfg=cfg, findings=[], score={"score": None}, conf={},
        endpoints=[], roadmap_items=[], exec_content=_make_exec_content(),
    )
    html = open(path, encoding="utf-8").read()

    assert "Acme Client" in html
    assert "Q3 2026 Readiness Assessment" in html
    assert "Jane Consultant" in html
    assert "2026-09-11" in html
    assert "STRICTLY CONFIDENTIAL" in html
    assert 'data:image/png;base64,' in html, "logo not embedded as a data URI"


def test_docx_full_branding_all_fields_present(tmp_path):
    from docx import Document

    from quirk.reports.docx_renderer import render_docx_report

    logo_path = tmp_path / "logo.png"
    logo_path.write_bytes(_png_bytes(255, 0, 0))

    cfg = _make_minimal_cfg(report=_report_cfg(
        logo_path=str(logo_path),
        client_name="Acme Client",
        engagement_name="Q3 2026 Readiness Assessment",
        prepared_by="Jane Consultant",
        cover_date="2026-09-11",
        confidentiality_line="STRICTLY CONFIDENTIAL — DO NOT DISTRIBUTE",
    ))

    path = str(tmp_path / "report.docx")
    result = render_docx_report(path=path, cfg=cfg, findings=[], exec_content=_make_exec_content())
    assert result is True

    doc = Document(path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    full_text += "\n" + doc.sections[0].header.paragraphs[0].text
    full_text += "\n" + doc.sections[0].footer.paragraphs[0].text

    for value in (
        "Acme Client",
        "Q3 2026 Readiness Assessment",
        "Jane Consultant",
        "2026-09-11",
        "STRICTLY CONFIDENTIAL",
    ):
        assert value in full_text, f"{value!r} missing from DOCX output"

    assert len(doc.inline_shapes) > 0, "logo picture not embedded in DOCX output"


# ---------------------------------------------------------------------------
# Partial branding — per-field conditionality
# ---------------------------------------------------------------------------


def test_html_partial_branding_only_set_field_present_absent_fields_missing(tmp_path):
    from quirk.reports.html_renderer import render_html_report

    cfg = _make_minimal_cfg(report=_report_cfg(client_name="Only Client Set"))

    path = str(tmp_path / "report.html")
    render_html_report(
        path=path, cfg=cfg, findings=[], score={"score": None}, conf={},
        endpoints=[], roadmap_items=[], exec_content=None,
    )
    html = open(path, encoding="utf-8").read()

    assert "Only Client Set" in html
    # The five unset fields must not leak any marker text — proving
    # per-field conditionality, not just presence of the one set field.
    assert "Engagement</span>" not in html
    assert "Prepared By</span>" not in html
    assert "Cover Date</span>" not in html
    assert "Confidentiality</span>" not in html
    assert 'data:image/png;base64,' not in html, "no logo path was set — must not embed one"


# ---------------------------------------------------------------------------
# Absent report section — byte-identical to a same-run baseline capture
# ---------------------------------------------------------------------------


def test_absent_report_section_renders_identically_to_captured_baseline(tmp_path):
    """A cfg with no `report` attribute renders the SAME as a cfg whose
    `report.branding` is explicitly all-None — proving "absent field = today's
    rendering" empirically (in-run comparison) rather than asserting intent."""
    from quirk.reports.html_renderer import render_html_report

    cfg_no_report_attr = SimpleNamespace(
        assessment=SimpleNamespace(
            name="Test Org", report_owner="Test Owner",
            data_classification="CONFIDENTIAL", timezone="UTC", logo_path=None,
        ),
        output=SimpleNamespace(directory=str(tmp_path)),
    )
    cfg_explicit_none_branding = _make_minimal_cfg(report=_report_cfg())

    path_a = str(tmp_path / "no_report_attr.html")
    path_b = str(tmp_path / "explicit_none_branding.html")
    render_html_report(
        path=path_a, cfg=cfg_no_report_attr, findings=[], score={"score": None},
        conf={}, endpoints=[], roadmap_items=[], exec_content=None,
    )
    render_html_report(
        path=path_b, cfg=cfg_explicit_none_branding, findings=[], score={"score": None},
        conf={}, endpoints=[], roadmap_items=[], exec_content=None,
    )

    html_a = open(path_a, encoding="utf-8").read()
    html_b = open(path_b, encoding="utf-8").read()

    # generated_at timestamps can differ by a second across the two calls —
    # normalize before the byte-identical comparison.
    import re
    norm = lambda s: re.sub(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2} UTC", "TS", s)
    assert norm(html_a) == norm(html_b), "no-report-attr cfg diverged from explicit-all-None-branding cfg"

    for marker in ("Client</span>", "Engagement</span>", "Prepared By</span>",
                   "Cover Date</span>", "Confidentiality</span>"):
        assert marker not in html_a
        assert marker not in html_b


# ---------------------------------------------------------------------------
# Logo precedence
# ---------------------------------------------------------------------------


def test_logo_precedence_report_branding_wins_over_assessment(tmp_path):
    import base64

    from quirk.reports.html_renderer import render_html_report

    report_logo = tmp_path / "report_logo.png"
    assessment_logo = tmp_path / "assessment_logo.png"
    report_logo_bytes = _png_bytes(255, 0, 0)
    assessment_logo_bytes = _png_bytes(0, 0, 255)
    report_logo.write_bytes(report_logo_bytes)
    assessment_logo.write_bytes(assessment_logo_bytes)

    cfg = _make_minimal_cfg(report=_report_cfg(logo_path=str(report_logo)))
    cfg.assessment.logo_path = str(assessment_logo)

    path = str(tmp_path / "report.html")
    render_html_report(
        path=path, cfg=cfg, findings=[], score={"score": None}, conf={},
        endpoints=[], roadmap_items=[], exec_content=None,
    )
    html = open(path, encoding="utf-8").read()

    report_b64 = base64.b64encode(report_logo_bytes).decode("ascii")
    assessment_b64 = base64.b64encode(assessment_logo_bytes).decode("ascii")
    assert report_b64 in html, "report.branding.logo_path should win when both are set"
    assert assessment_b64 not in html, "assessment.logo_path leaked when report.branding.logo_path was set"


def test_logo_precedence_assessment_logo_used_alone(tmp_path):
    """Phase 100 behaviour preserved: assessment.logo_path alone still embeds."""
    import base64

    from quirk.reports.html_renderer import render_html_report

    assessment_logo = tmp_path / "assessment_logo.png"
    assessment_logo_bytes = _png_bytes(0, 255, 0)
    assessment_logo.write_bytes(assessment_logo_bytes)

    cfg = _make_minimal_cfg(report=None)
    cfg.assessment.logo_path = str(assessment_logo)
    # No report.branding.logo_path at all — report is explicitly pinned None.

    path = str(tmp_path / "report.html")
    render_html_report(
        path=path, cfg=cfg, findings=[], score={"score": None}, conf={},
        endpoints=[], roadmap_items=[], exec_content=None,
    )
    html = open(path, encoding="utf-8").read()

    assessment_b64 = base64.b64encode(assessment_logo_bytes).decode("ascii")
    assert assessment_b64 in html


# ---------------------------------------------------------------------------
# Degradation — unreadable/nonexistent logo path
# ---------------------------------------------------------------------------


def test_degradation_nonexistent_logo_path_no_crash_either_surface(tmp_path):
    from docx import Document

    from quirk.reports.docx_renderer import render_docx_report
    from quirk.reports.html_renderer import render_html_report

    cfg = _make_minimal_cfg(report=_report_cfg(logo_path=str(tmp_path / "does-not-exist.png")))

    html_path = str(tmp_path / "report.html")
    render_html_report(
        path=html_path, cfg=cfg, findings=[], score={"score": None}, conf={},
        endpoints=[], roadmap_items=[], exec_content=None,
    )
    html = open(html_path, encoding="utf-8").read()
    assert 'data:image/png;base64,' not in html, "no image data URI expected for an unreadable logo"

    docx_path = str(tmp_path / "report.docx")
    result = render_docx_report(path=docx_path, cfg=cfg, findings=[])
    assert result is True, "DOCX render must complete despite an unreadable logo path"
    doc = Document(docx_path)
    assert len(doc.inline_shapes) == 0, "no picture expected for an unreadable logo"
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "[ Insert organization logo here ]" in full_text, "placeholder must be kept on degradation"


# ---------------------------------------------------------------------------
# Escaping — HTML-special branding values never appear raw
# ---------------------------------------------------------------------------


def test_html_escaping_of_html_special_branding_value(tmp_path):
    from quirk.reports.html_renderer import render_html_report

    cfg = _make_minimal_cfg(report=_report_cfg(client_name="Acme <b>&</b> Co"))

    path = str(tmp_path / "report.html")
    render_html_report(
        path=path, cfg=cfg, findings=[], score={"score": None}, conf={},
        endpoints=[], roadmap_items=[], exec_content=None,
    )
    html = open(path, encoding="utf-8").read()

    assert "<b>&</b>" not in html, "HTML-special branding value rendered raw"
    assert "Acme" in html and "Co" in html
