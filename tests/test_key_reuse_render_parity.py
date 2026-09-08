"""Phase 191 Plan 05 (SPKI-02 / D-01) — three-surface parity gate for the

"Key Reuse" section across CLI technical markdown, HTML, and DOCX.

Per-renderer caption duplication (`KEY_REUSE_ADVISORY_CAPTION` defined
independently in `quirk/reports/technical.py`, `quirk/reports/html_renderer.py`,
and `quirk/reports/docx_renderer.py`) is DELIBERATE — the established Phase 161
convention (see `tests/test_burndown_render_sections.py` and
`tests/test_vendor_trend_render_sections.py` for the direct analogs). This
file is the mechanism that keeps the three independently-worded copies from
drifting apart: a one-character edit to any of the three constants must make
`test_advisory_caption_is_identical_across_all_three_surfaces` fail.

This is a dedicated file, not an extension of `tests/test_cross_surface_parity.py`
(that file's scope is Phase 98 EXEC-04's narrative / top-risks parity — folding
a key-reuse assertion into it would widen a gate that is not about this
feature and blur its failure message).
"""
from __future__ import annotations

from types import SimpleNamespace

import pytest

# Auto-allowed by the skip-registry gate for a declared optional extra
# (pyproject.toml `[docx]` extra) — matching tests/test_cross_surface_parity.py's
# `test_docx_narrative_parity` convention. Gates every DOCX-leg test below.
pytest.importorskip("docx")

# ---------------------------------------------------------------------------
# Shared fixtures
# ---------------------------------------------------------------------------


def _key_reuse(**overrides) -> dict:
    """Synthetic `compute_key_reuse_clusters()` shape: one 4-member cluster,
    one 2-member cluster, already member-count descending (D-04)."""
    base = {
        "clusters": [
            {
                "fingerprint": "a" * 64,
                "member_count": 4,
                "cert_subject": "CN=shared-cert-a.example.com",
                "cert_pubkey_alg": "RSA",
                "cert_pubkey_size": 2048,
                "members": [
                    {"host": "host1.example.com", "port": 443},
                    {"host": "host2.example.com", "port": 443},
                    {"host": "host3.example.com", "port": 8443},
                    {"host": "host4.example.com", "port": 443},
                ],
            },
            {
                "fingerprint": "b" * 64,
                "member_count": 2,
                "cert_subject": "CN=shared-cert-b.example.com",
                "cert_pubkey_alg": "EC",
                "cert_pubkey_size": 256,
                "members": [
                    {"host": "host5.example.com", "port": 443},
                    {"host": "host6.example.com", "port": 443},
                ],
            },
        ],
        "fingerprinted": 9,
        "total": 12,
    }
    base.update(overrides)
    return base


def _empty_key_reuse() -> dict:
    return {"clusters": [], "fingerprinted": 9, "total": 12}


def _make_minimal_cfg(tmpdir="/tmp/quirk_test_key_reuse_parity"):
    return SimpleNamespace(
        assessment=SimpleNamespace(
            name="Key Reuse Parity Test Org",
            report_owner="Parity Tester",
            data_classification="CONFIDENTIAL",
            timezone="UTC",
            logo_path=None,
        ),
        output=SimpleNamespace(directory=tmpdir),
    )


def _exec_content(key_reuse=None):
    from quirk.reports.content_model import ExecContent

    return ExecContent(
        narrative_lead="Test narrative lead.",
        narrative_drivers=[],
        top_risks=[],
        roadmap_items=[],
        score_total=70,
        score_band="FAIR",
        subscores={},
        raw_sum=0,
        sev_counts={},
        key_reuse=key_reuse if key_reuse is not None else {},
    )


def _all_paragraph_texts(doc):
    return [p.text for p in doc.paragraphs]


def _all_table_texts(doc):
    cells = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cells.append(cell.text)
    return cells


def _docx_full_text(doc) -> str:
    return "\n".join(_all_paragraph_texts(doc) + _all_table_texts(doc))


# ---------------------------------------------------------------------------
# Caption parity — the load-bearing test
# ---------------------------------------------------------------------------


def test_advisory_caption_is_identical_across_all_three_surfaces():
    """Mirrors test_burndown_render_sections.py's node of the same name exactly.
    SPKI-02: three surfaces with three independently-worded advisory captions
    lets a client be shown a weaker caveat in whichever format they happen to
    read.
    """
    import inspect

    from quirk.reports import technical
    from quirk.reports.docx_renderer import KEY_REUSE_ADVISORY_CAPTION as docx_caption
    from quirk.reports.html_renderer import KEY_REUSE_ADVISORY_CAPTION as html_caption

    assert html_caption == docx_caption, (
        "SPKI-02: the HTML and DOCX key-reuse advisory captions have drifted apart"
    )
    assert html_caption in inspect.getsource(technical), (
        "SPKI-02: the CLI report's key-reuse caption no longer matches HTML/DOCX"
    )
    assert "Advisory" in html_caption


# ---------------------------------------------------------------------------
# Content parity — one synthetic payload, all three surfaces
# ---------------------------------------------------------------------------


def test_content_parity_across_all_three_surfaces(tmp_path):
    from docx import Document

    from quirk.reports.docx_renderer import render_docx_report
    from quirk.reports.html_renderer import render_key_reuse_section
    from quirk.reports.technical import build_tech_markdown

    payload = _key_reuse()

    html = render_key_reuse_section(payload)
    md = build_tech_markdown(_make_minimal_cfg(), [], [], key_reuse=payload)

    path = str(tmp_path / "key_reuse_parity.docx")
    render_docx_report(
        path=path,
        cfg=_make_minimal_cfg(str(tmp_path)),
        findings=[],
        exec_content=_exec_content(key_reuse=payload),
    )
    docx_text = _docx_full_text(Document(path))

    for surface_name, text in (("HTML", html), ("CLI markdown", md), ("DOCX", docx_text)):
        assert "9 of 12" in text, f"{surface_name} missing coverage numbers"
        assert "aaaaaaaaaaaaaaaa" in text, f"{surface_name} missing truncated fingerprint a"
        assert "bbbbbbbbbbbbbbbb" in text, f"{surface_name} missing truncated fingerprint b"
        assert "4" in text, f"{surface_name} missing 4-member cluster count"
        assert "2" in text, f"{surface_name} missing 2-member cluster count"

    # Biggest-cluster-first ordering (D-04) on every surface.
    for surface_name, text in (("HTML", html), ("CLI markdown", md), ("DOCX", docx_text)):
        idx_a = text.find("shared-cert-a.example.com")
        idx_b = text.find("shared-cert-b.example.com")
        assert idx_a != -1 and idx_b != -1, f"{surface_name} missing one of the clusters"
        assert idx_a < idx_b, f"{surface_name} did not order the 4-member cluster first"


def test_zero_reuse_sentence_identical_across_all_three_surfaces(tmp_path):
    from docx import Document

    from quirk.reports.docx_renderer import render_docx_report
    from quirk.reports.html_renderer import render_key_reuse_section
    from quirk.reports.technical import build_tech_markdown

    payload = _empty_key_reuse()
    expected = "No shared keys detected across 9 fingerprinted endpoints."

    html = render_key_reuse_section(payload)
    md = build_tech_markdown(_make_minimal_cfg(), [], [], key_reuse=payload)

    path = str(tmp_path / "key_reuse_zero.docx")
    render_docx_report(
        path=path,
        cfg=_make_minimal_cfg(str(tmp_path)),
        findings=[],
        exec_content=_exec_content(key_reuse=payload),
    )
    docx_text = _docx_full_text(Document(path))

    assert expected in html
    assert expected in md
    assert expected in docx_text


# ---------------------------------------------------------------------------
# Task 1 — HTML-specific behaviors
# ---------------------------------------------------------------------------


class TestHtmlKeyReuseSection:
    def test_ordering_and_leverage_phrasing(self):
        from quirk.reports.html_renderer import render_key_reuse_section

        html = render_key_reuse_section(_key_reuse())
        assert "Key Reuse" in html
        assert "9 of 12" in html
        assert "remediates 4 endpoints" in html
        idx_4 = html.find("shared-cert-a.example.com")
        idx_2 = html.find("shared-cert-b.example.com")
        assert idx_4 < idx_2

    def test_zero_clusters_renders_full_section(self):
        from quirk.reports.html_renderer import render_key_reuse_section

        html = render_key_reuse_section(_empty_key_reuse())
        assert "Key Reuse" in html
        assert "9 of 12" in html
        assert "No shared keys detected across 9 fingerprinted endpoints." in html

    def test_xss_shaped_cert_subject_is_escaped(self):
        from quirk.reports.html_renderer import render_key_reuse_section

        payload = _key_reuse()
        payload["clusters"][0]["cert_subject"] = "<script>alert(1)</script>"
        html = render_key_reuse_section(payload)
        assert "<script>" not in html

    def test_falsy_payload_returns_empty_string(self):
        from quirk.reports.html_renderer import render_key_reuse_section

        assert render_key_reuse_section({}) == ""
        assert render_key_reuse_section(None) == ""

    def test_section_reaches_full_html_report(self, tmp_path):
        from quirk.reports.html_renderer import render_html_report

        cfg = _make_minimal_cfg(str(tmp_path))
        path = str(tmp_path / "key_reuse_report.html")
        render_html_report(
            path=path,
            cfg=cfg,
            endpoints=[],
            findings=[],
            score={"score": 70, "rating": "FAIR", "subscores": {}, "drivers": []},
            conf={"confidence": 0},
            roadmap_items=[],
            exec_content=_exec_content(key_reuse=_key_reuse()),
        )
        with open(path, encoding="utf-8") as f:
            content = f.read()
        assert "Key Reuse" in content
        assert "remediates 4 endpoints" in content


# ---------------------------------------------------------------------------
# Task 2 — DOCX-specific behaviors
# ---------------------------------------------------------------------------


class TestDocxKeyReuseSection:
    def test_two_clusters_render_with_biggest_first(self, tmp_path):
        from docx import Document

        from quirk.reports.docx_renderer import render_docx_report

        path = str(tmp_path / "key_reuse_two_clusters.docx")
        render_docx_report(
            path=path,
            cfg=_make_minimal_cfg(str(tmp_path)),
            findings=[],
            exec_content=_exec_content(key_reuse=_key_reuse()),
        )
        doc = Document(path)
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert "Key Reuse" in headings

        text = _docx_full_text(doc)
        assert "Advisory - key reuse does not affect the readiness score." in text
        assert "9 of 12" in text
        idx_4 = text.find("shared-cert-a.example.com")
        idx_2 = text.find("shared-cert-b.example.com")
        assert idx_4 != -1 and idx_2 != -1
        assert idx_4 < idx_2
        assert "remediates 4 endpoints" in text

    def test_zero_clusters_emits_sentence_not_table(self, tmp_path):
        from docx import Document

        from quirk.reports.docx_renderer import render_docx_report

        path = str(tmp_path / "key_reuse_zero_clusters.docx")
        render_docx_report(
            path=path,
            cfg=_make_minimal_cfg(str(tmp_path)),
            findings=[],
            exec_content=_exec_content(key_reuse=_empty_key_reuse()),
        )
        doc = Document(path)
        text = _docx_full_text(doc)
        assert "No shared keys detected across 9 fingerprinted endpoints." in text
        # No key-reuse table was added when clusters is empty.
        for table in doc.tables:
            header_texts = [c.text for c in table.rows[0].cells]
            assert header_texts != [
                "Shared Key", "Public Key", "Members", "Endpoints", "Leverage"
            ]

    def test_empty_key_reuse_dict_does_not_raise(self, tmp_path):
        from quirk.reports.docx_renderer import render_docx_report

        path = str(tmp_path / "key_reuse_empty_dict.docx")
        result = render_docx_report(
            path=path,
            cfg=_make_minimal_cfg(str(tmp_path)),
            findings=[],
            exec_content=_exec_content(key_reuse={}),
        )
        assert result is True
